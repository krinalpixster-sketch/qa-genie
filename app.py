from flask import Flask, render_template, request, redirect, url_for, send_file, flash
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from openpyxl import Workbook
from google import genai
from dotenv import load_dotenv

import io
import os
import sqlite3
import markdown
import re
import difflib
import random

load_dotenv()

# Load API key from GEMINI_API_KEY env variable
api_key = os.getenv("GEMINI_API_KEY")
client = genai.Client(api_key=api_key) if api_key else None

app = Flask(__name__)
app.secret_key = "qa_genie_secret_key"

DATABASE = os.path.join(os.path.dirname(__file__), "qa_genie.db")


# -----------------------------
# Database Setup
# -----------------------------
def get_db():
    db_path = DATABASE
    if os.environ.get("VERCEL") or not os.access(os.path.dirname(DATABASE), os.W_OK):
        tmp_db = os.path.join("/tmp", "qa_genie.db")
        if not os.path.exists(tmp_db) and os.path.exists(DATABASE):
            import shutil
            try:
                shutil.copy2(DATABASE, tmp_db)
            except Exception as e:
                print("Error copying db to /tmp:", e)
        if os.path.exists(tmp_db):
            db_path = tmp_db

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    try:
        with get_db() as conn:
            conn.execute('''
                CREATE TABLE IF NOT EXISTS reports (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    type TEXT NOT NULL,
                    feature TEXT NOT NULL,
                    description TEXT NOT NULL,
                    content TEXT,
                    created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
                )
            ''')
            conn.commit()
    except Exception as e:
        print("Init DB warning:", e)


init_db()


# -----------------------------
# Helper & Similarity Functions
# -----------------------------
def calculate_similarity(text1, text2):
    if not text1 or not text2:
        return 0.0
    lines1 = set(line.strip().lower() for line in text1.splitlines() if line.strip() and not line.strip().startswith('|---'))
    lines2 = set(line.strip().lower() for line in text2.splitlines() if line.strip() and not line.strip().startswith('|---'))

    if not lines1 or not lines2:
        return 0.0

    intersection = lines1.intersection(lines2)
    smaller_len = min(len(lines1), len(lines2))
    overlap_ratio = len(intersection) / smaller_len if smaller_len > 0 else 0.0

    matcher_ratio = difflib.SequenceMatcher(None, text1.lower(), text2.lower()).ratio()
    return max(overlap_ratio, matcher_ratio)


def normalize_markdown_table_ids(md_text, prefix="TC"):
    if not md_text:
        return md_text

    import re
    if prefix == "BUG" and ("**Bug ID:**" in md_text or "Bug ID:" in md_text):
        md_text = re.sub(r'(\*\*Bug ID:\*\*|\bBug ID:)\s*[^\n]+', r'\1 BUG001', md_text, count=1)

    lines = md_text.splitlines()
    new_lines = []
    idx = 1
    inside_table = False

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            parts = [cell.strip() for cell in stripped.split("|")[1:-1]]
            if not parts:
                new_lines.append(line)
                continue

            if all(set(cell) <= set("-: ") for cell in parts):
                new_lines.append(line)
                inside_table = True
                continue

            header_first = parts[0].lower()
            if not inside_table and ("id" in header_first or "case" in header_first or "scenario" in header_first or "bug" in header_first or "check" in header_first):
                new_lines.append(line)
                continue

            if inside_table:
                parts[0] = f"{prefix}{idx:03d}"
                idx += 1
                new_lines.append("| " + " | ".join(parts) + " |")
            else:
                new_lines.append(line)
        else:
            inside_table = False
            new_lines.append(line)

    return "\n".join(new_lines)


def parse_markdown_table_data(md_text):
    rows = []
    if not md_text:
        return rows
    lines = md_text.splitlines()
    for line in lines:
        line = line.strip()
        if line.startswith("|") and line.endswith("|"):
            parts = [cell.strip() for cell in line.split("|")[1:-1]]
            if not parts:
                continue
            if all(set(cell) <= set("-: ") for cell in parts):
                continue
            rows.append(parts)
    return rows


def parse_and_combine_markdown_tables(*texts):
    unique_cases = {}
    rows_list = []

    for text in texts:
        if not text:
            continue
        lines = text.splitlines()
        for line in lines:
            line = line.strip()
            if line.startswith("|") and line.endswith("|"):
                parts = [p.strip() for p in line.split("|")[1:-1]]
                if not parts:
                    continue
                header_check = parts[0].lower().replace("-", "")
                if header_check in ["test case id", "id", ""] or all(set(p) <= set("-: ") for p in parts):
                    continue

                key = parts[3].lower() if len(parts) >= 4 else "|".join(parts).lower()
                if key not in unique_cases:
                    unique_cases[key] = parts
                    rows_list.append(parts)

    formatted_rows = []
    for idx, parts in enumerate(rows_list, start=1):
        tc_id = f"TC{idx:03d}"
        if len(parts) >= 5:
            formatted_rows.append([tc_id, parts[1], parts[2], parts[3], parts[4]])
        elif len(parts) == 4:
            formatted_rows.append([tc_id, parts[0], parts[1], parts[2], parts[3]])
        else:
            formatted_rows.append([tc_id, "Functional", "Medium", parts[0], parts[-1] if len(parts) > 1 else "Success"])

    return formatted_rows


def render_markdown_table_from_rows(rows):
    md = "| Test Case ID | Test Type | Priority | Test Case | Expected Result |\n"
    md += "|---|---|---|---|---|\n"
    for r in rows:
        md += f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} | {r[4]} |\n"
    return md


# -----------------------------
# AI Generators & Dynamic Fallbacks
# -----------------------------
def get_fallback_test_cases(feature, description):
    rows = [
        ["TC_001", "Functional", "High", f"Verify user can access the {feature} feature.", "Feature page loads successfully."],
        ["TC_002", "Functional", "High", "Verify all mandatory fields are displayed.", "Mandatory fields are visible and marked."],
        ["TC_003", "Functional", "Medium", "Verify all optional fields are displayed.", "Optional fields are visible."],
        ["TC_004", "Functional", "High", "Verify valid input is accepted.", "Input accepted without error."],
        ["TC_005", "Functional", "High", "Verify invalid input shows validation message.", "Appropriate validation error message is shown."],
        ["TC_006", "Functional", "High", "Verify Save button works correctly.", "Data is saved and confirmation is displayed."],
        ["TC_007", "Functional", "Medium", "Verify Cancel button works correctly.", "Operation cancelled and form cleared/closed."],
        ["TC_008", "Functional", "Medium", "Verify success message is displayed after action.", "Success notification popup appears."],
        ["TC_009", "Functional", "High", "Verify data is saved successfully in database.", "Data persisted correctly in database."],
        ["TC_010", "Functional", "Medium", "Verify page refresh does not lose unsaved input.", "Draft data restored or warning shown."],
        ["TC_011", "Negative", "High", "Leave mandatory fields empty and submit.", "Validation errors highlight mandatory fields."],
        ["TC_012", "Negative", "High", "Enter invalid characters in text inputs.", "System blocks input or shows error message."],
        ["TC_013", "Negative", "Medium", "Enter text exceeding maximum length boundary.", "Input truncated or character limit alert shown."],
        ["TC_014", "Negative", "Medium", "Enter zero or negative numbers in numeric fields.", "Numeric range error message displayed."],
        ["TC_015", "Negative", "High", "Enter SQL injection string in input field.", "Input sanitized safely without syntax error."],
        ["TC_016", "Negative", "High", "Enter XSS payload script in text field.", "Script HTML encoded and rendered safely."],
        ["TC_017", "Negative", "Medium", "Submit request without network connection.", "Offline warning alert displayed."],
        ["TC_018", "Negative", "High", "Click Submit button rapidly multiple times.", "Form disabled after first click preventing duplicate submission."],
        ["TC_019", "Edge", "High", "Submit form with exact minimum character limits.", "Form submits successfully."],
        ["TC_020", "Edge", "High", "Submit form with exact maximum character limits.", "Form submits successfully."],
        ["TC_021", "Edge", "Medium", "Paste large block of text into input field.", "Pasted content handled smoothly."],
        ["TC_022", "Edge", "Low", "Use non-English Unicode characters in inputs.", "Unicode characters saved and displayed properly."],
        ["TC_023", "Edge", "Medium", "Access feature with session timeout active.", "User redirected to login screen."],
        ["TC_024", "Edge", "Medium", "Test concurrently on two browser tabs.", "Session state maintained synchronously."],
        ["TC_025", "Edge", "Low", "Test with slow 3G internet network speed.", "Loading spinner displayed until completion."],
        ["TC_026", "UI", "High", "Verify responsiveness on Mobile screens.", "Layout adjusts cleanly to mobile width."],
        ["TC_027", "UI", "High", "Verify responsiveness on Tablet screens.", "Layout scales cleanly to tablet screen."],
        ["TC_028", "UI", "Medium", "Verify font style, size and contrast ratio.", "Text meets accessibility contrast guidelines."],
        ["TC_029", "UI", "Medium", "Verify hover states and active button styles.", "Buttons provide visual feedback on interaction."],
        ["TC_030", "UI", "Low", "Verify alignment of form labels and inputs.", "Elements aligned neatly."],
        ["TC_031", "Regression", "High", "Verify user login state remains valid.", "User remains authenticated."],
        ["TC_032", "Regression", "High", "Verify navigation bar links operate normally.", "Navigation routing works."],
        ["TC_033", "Regression", "High", "Verify API endpoints return 200 OK status.", "API response successful."],
        ["TC_034", "Regression", "Medium", "Verify application performance under load.", "Page response time within SLA limit."],
        ["TC_035", "Regression", "Medium", "Verify logout clears session tokens.", "Session destroyed upon logout."]
    ]
    return render_markdown_table_from_rows(rows)


def get_fallback_test_cases_regenerated(feature, description):
    variation = random.choice([1, 2, 3])
    if variation == 1:
        rows = [
            ["TC_101", "Security", "High", f"Verify rate-limiting on {feature} submission endpoint against brute-force attacks.", "HTTP 429 Too Many Requests returned after 10 rapid attempts."],
            ["TC_102", "Security", "High", f"Verify SQL injection vulnerability on {feature} input fields.", "Input sanitized properly and no database syntax error is thrown."],
            ["TC_103", "Security", "High", f"Verify CSRF token protection during {feature} form submission.", "Request without valid CSRF token is rejected with 403 Forbidden."],
            ["TC_104", "Security", "Medium", f"Verify XSS payload injection in {feature} text fields.", "Input HTML encoded safely before rendering in DOM."],
            ["TC_105", "Security", "High", f"Verify sensitive parameters in {feature} URL query strings.", "No passwords, tokens, or PII exposed in URL query parameters."],
            ["TC_106", "Performance", "High", f"Verify response time of {feature} under 500 concurrent user requests.", "Response latency remains below 1.5 seconds."],
            ["TC_107", "Performance", "Medium", f"Verify CPU and RAM memory usage during bulk processing in {feature}.", "Memory consumption remains within server SLA limits."],
            ["TC_108", "Performance", "High", f"Verify database query optimization under 10,000 records in {feature}.", "Query execution uses indexed columns without full table scan."],
            ["TC_109", "Performance", "Medium", f"Verify HTTP Cache-Control headers on {feature} static responses.", "Private user data not cached by intermediate CDNs."],
            ["TC_110", "Integration", "High", f"Verify third-party webhook failure recovery during {feature} processing.", "System retries with exponential backoff algorithm."],
            ["TC_111", "Integration", "High", f"Verify API endpoint response payload schema for {feature}.", "JSON response matches OpenAPI spec with all required fields."],
            ["TC_112", "Integration", "Medium", f"Verify behavior of {feature} when external microservice drops connection.", "Circuit breaker opens and graceful fallback data is served."],
            ["TC_113", "UI/Accessibility", "Medium", f"Verify screen reader (VoiceOver/NVDA) accessibility on {feature} elements.", "ARIA labels and input focus are read clearly."],
            ["TC_114", "UI/Accessibility", "Low", f"Verify WCAG AA color contrast ratio on {feature} interactive buttons.", "Contrast ratio exceeds 4.5:1 requirement."],
            ["TC_115", "UI/Accessibility", "Medium", f"Verify tab keyboard navigation through all {feature} input fields.", "Focus ring moves logically in sequential order."],
            ["TC_116", "Functional", "High", f"Verify session state restoration after unexpected browser crash during {feature}.", "Draft inputs restored smoothly from local storage."],
            ["TC_117", "Functional", "High", f"Verify idempotency key handling during duplicate submission of {feature}.", "Only one transaction created; subsequent requests return cached result."],
            ["TC_118", "Functional", "Medium", f"Verify multi-tab concurrent interaction on {feature}.", "Session state synced synchronously across browser tabs."],
            ["TC_119", "Functional", "High", f"Verify rollback mechanism if error occurs mid-transaction in {feature}.", "Database state rolls back without partial orphan records."],
            ["TC_120", "Negative", "High", f"Verify invalid file payload upload to {feature}.", "System rejects corrupted payload with validation alert."],
            ["TC_121", "Negative", "High", f"Verify submission of empty payload JSON body to {feature}.", "HTTP 400 Bad Request returned with clear validation details."],
            ["TC_122", "Negative", "Medium", f"Verify submission of expired JWT authentication token to {feature}.", "HTTP 401 Unauthorized returned and user redirected to login."],
            ["TC_123", "Negative", "High", f"Verify parameter tampering on user ID field in {feature}.", "System enforces authorization check and blocks horizontal escalation."],
            ["TC_124", "Negative", "Medium", f"Verify submitting non-numeric string in integer fields of {feature}.", "Type conversion error caught cleanly without 500 crash."],
            ["TC_125", "Edge", "Medium", f"Verify {feature} behavior during database failover maintenance window.", "Graceful maintenance notice displayed."],
            ["TC_126", "Edge", "High", f"Verify submitting exact maximum boundary payload size to {feature}.", "Payload uploaded successfully without memory leak."],
            ["TC_127", "Edge", "Low", f"Verify system clock timezone offset changes during {feature} use.", "Timestamps recorded in standard UTC format."],
            ["TC_128", "Edge", "Medium", f"Verify non-English Unicode and emoji input handling in {feature}.", "Unicode characters saved and rendered without corruption."],
            ["TC_129", "Edge", "Low", f"Verify behavior of {feature} on slow 2G/3G mobile network connection.", "Loading indicator displayed until payload transmission finishes."],
            ["TC_130", "Regression", "High", f"Verify backwards compatibility of {feature} with legacy v1 API schema.", "JSON payload parses without missing field errors."],
            ["TC_131", "Regression", "High", f"Verify user authentication state validity after navigating {feature}.", "Authentication cookie remains secure and httpOnly."],
            ["TC_132", "Regression", "Medium", f"Verify audit log event recorded upon successful action in {feature}.", "Audit table records user ID, timestamp, and action detail."],
            ["TC_133", "Regression", "High", f"Verify database connection pool auto-recovers after network drop during {feature}.", "Connection pool restores without restart."],
            ["TC_134", "Regression", "Medium", f"Verify clear success confirmation Toast notification in {feature}.", "Toast notification auto-dismisses after 5 seconds."],
            ["TC_135", "Regression", "High", f"Verify logout revokes active refresh tokens associated with {feature}.", "Session tokens invalidated on auth server."]
        ]
    elif variation == 2:
        rows = [
            ["TC_201", "Authentication", "High", f"Verify active JWT token expiration during {feature} session.", "User redirected to re-authentication screen."],
            ["TC_202", "Authentication", "High", f"Verify multi-factor authentication (MFA) enforcement on {feature}.", "MFA challenge requested before action completes."],
            ["TC_203", "Concurrency", "High", f"Verify parallel form submission on two separate tabs for {feature}.", "Idempotency key prevents duplicate database insertion."],
            ["TC_204", "Concurrency", "High", f"Verify optimistic locking mechanism during simultaneous edits on {feature}.", "Version conflict error shown if data modified by another user."],
            ["TC_205", "Network", "Medium", f"Verify behavior of {feature} when switching between Wi-Fi and 4G.", "Request retried seamlessly without data loss."],
            ["TC_206", "Network", "High", f"Verify offline queueing mechanism for {feature} in Progressive Web App.", "Actions queued and synced upon reconnecting."],
            ["TC_207", "Localization", "Low", f"Verify multi-language UTF-8 character rendering in {feature}.", "Special characters display correctly without corruption."],
            ["TC_208", "Localization", "Medium", f"Verify right-to-left (RTL) language layout alignment for {feature}.", "UI flips horizontally with correct text direction."],
            ["TC_209", "Boundary", "High", f"Verify exact maximum file payload size upload to {feature}.", "Payload uploaded successfully without memory leak."],
            ["TC_210", "Boundary", "High", f"Verify zero length input handling on optional text inputs in {feature}.", "Empty input treated as NULL without string error."],
            ["TC_211", "Usability", "Medium", f"Verify tab keyboard navigation through all {feature} input fields.", "Focus ring moves logically in sequential order."],
            ["TC_212", "Usability", "Low", f"Verify hover states and tooltip explanations on {feature} icons.", "Tooltip displays contextual help on mouse hover."],
            ["TC_213", "Negative", "High", f"Verify malformed JSON payload handling in {feature} endpoint.", "HTTP 400 Bad Request returned with clear error message."],
            ["TC_214", "Negative", "High", f"Verify uploading executable file (.exe/.sh) to {feature} file upload.", "File extension whitelist blocks dangerous file types."],
            ["TC_215", "Performance", "High", f"Verify CPU and memory usage during bulk processing in {feature}.", "Memory consumption remains within SLA limits."],
            ["TC_216", "Performance", "Medium", f"Verify response compression (gzip/brotli) on {feature} assets.", "Content-Encoding header indicates brotli/gzip compression."],
            ["TC_217", "Compatibility", "Medium", f"Verify rendering of {feature} on Safari Mobile and Chrome Android.", "Layout scales responsively without horizontal scroll."],
            ["TC_218", "Compatibility", "Low", f"Verify rendering of {feature} on legacy Edge and Firefox ESR.", "All essential controls display and function properly."],
            ["TC_219", "Security", "High", f"Verify XSS payload injection in {feature} text fields.", "Input HTML encoded safely before DOM rendering."],
            ["TC_220", "Security", "High", f"Verify CORS headers on {feature} cross-origin request.", "Access-Control-Allow-Origin restricted to trusted domains."],
            ["TC_221", "Functional", "High", f"Verify data persistence after saving {feature} form.", "Record appears immediately in summary list."],
            ["TC_222", "Functional", "High", f"Verify field clearing when clicking Reset on {feature} form.", "All form inputs cleared to default values."],
            ["TC_223", "Functional", "Medium", f"Verify default values pre-populated on {feature} page load.", "Default settings loaded accurately."],
            ["TC_224", "Functional", "High", f"Verify error boundary handles unexpected React/Vue component crash in {feature}.", "Fallback UI displayed instead of white screen."],
            ["TC_225", "Edge", "Medium", f"Verify submitting maximum permitted multi-select options in {feature}.", "All selected options saved successfully."],
            ["TC_226", "Edge", "High", f"Verify session timeout handling while submitting {feature}.", "Draft saved and user prompted to re-login."],
            ["TC_227", "Edge", "Low", f"Verify pasting formatted rich text into plain text input of {feature}.", "Formatting stripped cleanly, keeping plain text."],
            ["TC_228", "Regression", "High", f"Verify database transaction isolation level during {feature} updates.", "Read Committed isolation prevents dirty reads."],
            ["TC_229", "Regression", "High", f"Verify API authentication middleware executes on all {feature} sub-routes.", "Unauthenticated requests blocked at middleware layer."],
            ["TC_230", "Regression", "Medium", f"Verify automated email notification dispatched upon {feature} completion.", "Email delivered with correct template variables."],
            ["TC_231", "Regression", "Medium", f"Verify pagination controls when list size exceeds 50 items in {feature}.", "Next/Previous buttons navigate through pages."],
            ["TC_232", "Regression", "Low", f"Verify print stylesheet layout for {feature} report.", "Navigation UI hidden in print preview."],
            ["TC_233", "Security", "High", f"Verify clickjacking defense on {feature} pages.", "X-Frame-Options set to DENY or SAMEORIGIN."],
            ["TC_234", "Security", "High", f"Verify HTTPS redirection enforcement for {feature}.", "HTTP traffic auto-redirected to HTTPS."],
            ["TC_235", "Regression", "High", f"Verify audit trail logs IP address and User-Agent for {feature} changes.", "Audit record contains client metadata."]
        ]
    else:
        rows = [
            ["TC_301", "Data Privacy", "High", f"Verify GDPR compliance data masking on {feature} output.", "Sensitive PII data masked with asterisks."],
            ["TC_302", "Data Privacy", "High", f"Verify data retention policy auto-purges old {feature} records.", "Records older than retention window deleted."],
            ["TC_303", "API Security", "High", f"Verify OAuth2 scope permissions for {feature} endpoint access.", "Unauthorized scope returns 403 Forbidden."],
            ["TC_304", "API Security", "High", f"Verify API key rotation handling for {feature} integration.", "New key accepted; old key revoked gracefully."],
            ["TC_305", "Caching", "Medium", f"Verify HTTP Cache-Control headers on {feature} responses.", "Private data not stored in public CDN cache."],
            ["TC_306", "Caching", "Medium", f"Verify ETag validation headers for {feature} resource requests.", "HTTP 304 Not Modified returned for unchanged content."],
            ["TC_307", "Resilience", "High", f"Verify circuit breaker pattern during downstream dependency failure in {feature}.", "Fallback data returned without crashing."],
            ["TC_308", "Resilience", "High", f"Verify bulk failure retry queue for {feature} background jobs.", "Failed jobs moved to dead-letter queue."],
            ["TC_309", "Usability", "Medium", f"Verify clear success confirmation Toast notification in {feature}.", "Toast notification auto-dismisses after 5 seconds."],
            ["TC_310", "Usability", "Medium", f"Verify loading skeleton placeholders displayed while loading {feature}.", "Skeleton layout prevents visual layout shift."],
            ["TC_311", "Negative", "High", f"Submit expired session token to {feature}.", "Session cleared and login page shown."],
            ["TC_312", "Negative", "High", f"Submit negative currency values in {feature} transaction fields.", "Validation error prevents negative financial calculation."],
            ["TC_313", "Negative", "Medium", f"Submit invalid date format (e.g. 31/02/2026) to {feature}.", "Date parser rejects invalid calendar dates."],
            ["TC_314", "Edge", "Low", f"Verify system clock timezone offset changes during {feature} use.", "Timestamps recorded in standard UTC."],
            ["TC_315", "Edge", "Medium", f"Verify high screen resolution scaling (4K/Retina) for {feature}.", "Graphics and text remain crisp without pixelation."],
            ["TC_316", "Regression", "High", f"Verify database index performance on {feature} search query.", "Query execution plan utilizes index scan."],
            ["TC_317", "Regression", "High", f"Verify cascading delete constraints for {feature} entity relationships.", "Child records cleaned up without foreign key orphan."],
            ["TC_318", "Mobile", "Medium", f"Verify touch screen gesture interactions for {feature} on mobile.", "Tap and swipe actions respond fluidly."],
            ["TC_319", "Mobile", "Medium", f"Verify virtual keyboard auto-correction settings on {feature} inputs.", "Email inputs set to autocomplete='email'."],
            ["TC_320", "Stress", "High", f"Verify 10,000 rapid requests burst load on {feature}.", "System throttles gracefully without dropping database connection."],
            ["TC_321", "Functional", "High", f"Verify filtering results by category in {feature}.", "Grid updates displaying matching records."],
            ["TC_322", "Functional", "Medium", f"Verify sorting columns in ascending and descending order for {feature}.", "Data reordered accurately."],
            ["TC_323", "Functional", "High", f"Verify exporting {feature} summary to CSV/Excel.", "Exported file contains matching dataset."],
            ["TC_324", "UI", "Low", f"Verify responsive breakpoint behavior at 768px for {feature}.", "Navigation drawer collapses into hamburger menu."],
            ["TC_325", "UI", "Medium", f"Verify modal dialog traps keyboard focus within {feature}.", "Tab key cycles inside open modal."],
            ["TC_326", "Security", "High", f"Verify HTTP Strict Transport Security (HSTS) header on {feature}.", "Strict-Transport-Security header present in response."],
            ["TC_327", "Security", "High", f"Verify secure cookie flags (Secure, HttpOnly, SameSite=Lax) on {feature}.", "Cookie flags set correctly."],
            ["TC_328", "Negative", "High", f"Bypass required field validation via API tool on {feature}.", "Server-side validator rejects invalid request."],
            ["TC_329", "Edge", "Medium", f"Submit input containing zero-width space characters to {feature}.", "Sanitizer trims zero-width spaces."],
            ["TC_330", "Integration", "High", f"Verify queue message consumption rate for {feature}.", "RabbitMQ/Kafka consumer keeps pace with producer."],
            ["TC_331", "Regression", "High", f"Verify user role permission changes take effect immediately in {feature}.", "Permission revoked user blocked on next API call."],
            ["TC_332", "Regression", "Medium", f"Verify error logging captures stack trace for {feature} errors.", "Sentry/LogRocket records error traceback."],
            ["TC_333", "Usability", "Low", f"Verify clear inline error indicators under failing {feature} fields.", "Red border and error text highlight field."],
            ["TC_334", "Performance", "Medium", f"Verify Asset bundle size for {feature} JavaScript module.", "Module chunk size under 200KB gzipped."],
            ["TC_335", "Regression", "High", f"Verify database connection pool leak check during {feature} load test.", "Zero connection leaks detected after test run."]
        ]
    return render_markdown_table_from_rows(rows)


def generate_ai_test_cases(feature, description):
    if not client:
        return get_fallback_test_cases(feature, description)

    prompt1 = f"""
    Generate 20 unique QA test cases for the following feature.
    Feature: {feature}
    Description: {description}

    Return ONLY a Markdown table with exact columns:
    | Test Case ID | Test Type | Priority | Test Case | Expected Result |
    """

    prompt2 = f"""
    Generate 20 additional unique QA test cases for the following feature.
    Feature: {feature}
    Description: {description}

    Generate completely different test cases. Do not repeat any test case from the first batch.

    Return ONLY a Markdown table with exact columns:
    | Test Case ID | Test Type | Priority | Test Case | Expected Result |
    """

    try:
        print("AI GENERATION STARTED - BATCH 1")
        resp1 = client.models.generate_content(model="gemini-3.5-flash", contents=prompt1)
        text1 = resp1.text or ""

        print("AI GENERATION STARTED - BATCH 2")
        resp2 = client.models.generate_content(model="gemini-3.5-flash", contents=prompt2)
        text2 = resp2.text or ""

        combined_rows = parse_and_combine_markdown_tables(text1, text2)
        print(f"AI RESPONSE RECEIVED - Total unique test cases: {len(combined_rows)}")

        if len(combined_rows) >= 15:
            return render_markdown_table_from_rows(combined_rows)
        else:
            return get_fallback_test_cases(feature, description)
    except Exception as e:
        print("AI GENERATION FAILED:", e)
        return get_fallback_test_cases(feature, description)


def generate_ai_test_cases_regenerated(feature, description, previous_content):
    if not client:
        return get_fallback_test_cases_regenerated(feature, description)

    prompt1 = f"""
    Generate 20 completely NEW and unique QA test cases for the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Generate a completely new set of test cases with fresh QA coverage.
    - Do NOT repeat any test case from the previous result provided below.
    - Do NOT merely rephrase or reorder previous test cases.
    - Explore different functional flows, security angles, performance parameters, edge cases, mobile UX, and regression paths.

    PREVIOUS RESULT CONTENT TO AVOID:
    {previous_content[:1500]}

    Return ONLY a Markdown table with exact columns:
    | Test Case ID | Test Type | Priority | Test Case | Expected Result |
    """

    prompt2 = f"""
    Generate 20 ADDITIONAL unique QA test cases for the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Do NOT repeat any test case from batch 1 or the previous result below.

    PREVIOUS RESULT CONTENT TO AVOID:
    {previous_content[:1500]}

    Return ONLY a Markdown table with exact columns:
    | Test Case ID | Test Type | Priority | Test Case | Expected Result |
    """

    try:
        print("AI REGENERATION STARTED - BATCH 1")
        resp1 = client.models.generate_content(model="gemini-3.5-flash", contents=prompt1)
        text1 = resp1.text or ""

        print("AI REGENERATION STARTED - BATCH 2")
        resp2 = client.models.generate_content(model="gemini-3.5-flash", contents=prompt2)
        text2 = resp2.text or ""

        combined_rows = parse_and_combine_markdown_tables(text1, text2)
        if len(combined_rows) >= 15:
            return render_markdown_table_from_rows(combined_rows)
    except Exception as e:
        print("AI TEST CASE REGENERATION FAILED:", e)

    return get_fallback_test_cases_regenerated(feature, description)


def get_fallback_scenarios(feature, description):
    rows = [
        ["TS001", f"Verify complete end-to-end happy path workflow for {feature}."],
        ["TS002", f"Verify {feature} behavior when submitting valid mandatory field data."],
        ["TS003", f"Verify {feature} behavior when submitting optional field data."],
        ["TS004", f"Verify form validation messages when mandatory fields are left blank in {feature}."],
        ["TS005", f"Verify input field restrictions when entering special characters in {feature}."],
        ["TS006", f"Verify max character length boundaries for text input fields in {feature}."],
        ["TS007", f"Verify role-based access control and unauthorized access prevention for {feature}."],
        ["TS008", f"Verify system behavior when session expires while interacting with {feature}."],
        ["TS009", f"Verify error handling and recovery when server returns HTTP 500 error in {feature}."],
        ["TS010", f"Verify system behavior under network disconnect and auto-reconnect retry during {feature} operation."],
        ["TS011", f"Verify duplicate form submission prevention upon rapid multiple clicks in {feature}."],
        ["TS012", f"Verify data persistence after saving records in {feature}."],
        ["TS013", f"Verify browser back and forward button navigation during multi-step {feature} flow."],
        ["TS014", f"Verify page responsiveness and UI layout alignment across mobile and tablet viewports for {feature}."],
        ["TS015", f"Verify keyboard tab navigation and focus indicator highlights across {feature} input fields."],
        ["TS016", f"Verify XSS payload sanitization on text input fields in {feature}."],
        ["TS017", f"Verify SQL injection security prevention on query parameters in {feature}."],
        ["TS018", f"Verify performance response latency under average concurrent user load for {feature}."],
        ["TS019", f"Verify search and filtering results by category in {feature}."],
        ["TS020", f"Verify column sorting in ascending and descending order for {feature} dataset."],
        ["TS021", f"Verify data export summary to PDF and Excel for {feature}."],
        ["TS022", f"Verify clear success Toast notification feedback upon completing action in {feature}."],
        ["TS023", f"Verify audit trail logging records user ID, timestamp, and action details for {feature}."],
        ["TS024", f"Verify graceful failure handling when external API dependency drops during {feature} transaction."],
        ["TS025", f"Verify session token invalidation upon user logout from {feature}."]
    ]
    md = "| Scenario ID | Test Scenario |\n|---|---|\n"
    for r in rows:
        md += f"| {r[0]} | {r[1]} |\n"
    return md


def get_fallback_scenarios_regenerated(feature, description):
    v = random.choice([1, 2, 3])
    if v == 1:
        rows = [
            ["TS001", f"Verify multi-factor authentication (MFA) step-up challenge during high-risk action in {feature}."],
            ["TS002", f"Verify webhook notification event triggers after state change in {feature}."],
            ["TS003", f"Verify offline mode behavior and sync queuing for {feature} in Progressive Web App mode."],
            ["TS004", f"Verify UI responsiveness and touch gesture interactions on iOS Safari and Android Chrome for {feature}."],
            ["TS005", f"Verify bulk batch processing workflow for {feature} with 1,000 records."],
            ["TS006", f"Verify database transaction rollback on downstream service timeout in {feature}."],
            ["TS007", f"Verify double-clicking submit control repeatedly during network latency in {feature}."],
            ["TS008", f"Verify audit trail logging records user ID, IP address, and timestamp for {feature}."],
            ["TS009", f"Verify parameter values at exact maximum capacity threshold for {feature}."],
            ["TS010", f"Verify system recovery after force-killing application process during {feature} save."],
            ["TS011", f"Verify zero-width space character trimming from text input fields in {feature}."],
            ["TS012", f"Verify rate-limiting on {feature} endpoint against brute-force automated requests."],
            ["TS013", f"Verify optimistic locking version conflict handling during simultaneous edits in {feature}."],
            ["TS014", f"Verify memory consumption remains within SLA limits during prolonged use of {feature}."],
            ["TS015", f"Verify dark mode theme rendering and color contrast ratios across {feature} UI elements."],
            ["TS016", f"Verify automatic session renewal when active JWT bearer token approaches expiry during {feature}."],
            ["TS017", f"Verify screen reader ARIA landmarks and field voice labels across {feature} controls."],
            ["TS018", f"Verify HTTP Cache-Control headers prevent intermediate CDN caching of private data in {feature}."],
            ["TS019", f"Verify cascading delete constraints clean up associated entity relationships in {feature}."],
            ["TS020", f"Verify database index utilization for high-speed search queries in {feature}."],
            ["TS021", f"Verify draft restoration from local storage after unexpected browser tab crash in {feature}."],
            ["TS022", f"Verify multi-tab concurrent session synchronization across browser tabs for {feature}."],
            ["TS023", f"Verify file upload whitelist blocks executable payload files in {feature}."],
            ["TS024", f"Verify CORS headers restrict unauthorized cross-origin requests to {feature} API."],
            ["TS025", f"Verify automatic HTTPS redirection enforcement for all {feature} HTTP endpoints."]
        ]
    elif v == 2:
        rows = [
            ["TS001", f"Verify parallel form submissions on separate browser tabs for {feature}."],
            ["TS002", f"Verify API rate limiting behavior under burst request traffic on {feature}."],
            ["TS003", f"Verify GDPR data privacy masking on sensitive output fields in {feature}."],
            ["TS004", f"Verify keyboard accessibility navigation cycling through all controls in {feature}."],
            ["TS005", f"Verify REST API backwards compatibility with legacy client payload schemas in {feature}."],
            ["TS006", f"Verify CSRF token validation and XSS input sanitization in {feature}."],
            ["TS007", f"Verify DOM node virtual scrolling responsiveness when dataset exceeds 10,000 rows in {feature}."],
            ["TS008", f"Verify error handling when invalid OAuth authorization headers are passed to {feature}."],
            ["TS009", f"Verify zero-length and whitespace-only input handling in {feature}."],
            ["TS010", f"Verify automatic data purge policy for archived records older than retention window in {feature}."],
            ["TS011", f"Verify API key rotation handling without service interruption for {feature}."],
            ["TS012", f"Verify ETag validation headers return HTTP 304 Not Modified for unchanged data in {feature}."],
            ["TS013", f"Verify circuit breaker pattern during downstream microservice outage in {feature}."],
            ["TS014", f"Verify dead-letter queue routing for failed background jobs in {feature}."],
            ["TS015", f"Verify loading skeleton placeholders prevent layout shift while loading {feature}."],
            ["TS016", f"Verify validation error prevention for negative currency values in {feature}."],
            ["TS017", f"Verify date parser rejects invalid calendar date strings in {feature}."],
            ["TS018", f"Verify system clock timezone offset changes during active session in {feature}."],
            ["TS019", f"Verify high-resolution Retina display graphics rendering for {feature}."],
            ["TS020", f"Verify virtual keyboard auto-correction and auto-complete hints on mobile inputs in {feature}."],
            ["TS021", f"Verify burst load handling of 5,000 concurrent requests without connection leaks in {feature}."],
            ["TS022", f"Verify modal dialog traps keyboard focus within open modal for {feature}."],
            ["TS023", f"Verify HSTS Strict Transport Security headers on response payload in {feature}."],
            ["TS024", f"Verify cookie security flags Secure, HttpOnly, and SameSite=Lax on {feature} cookies."],
            ["TS025", f"Verify clear inline error indicators under failing form fields in {feature}."]
        ]
    else:
        rows = [
            ["TS001", f"Verify OAuth2 scope authorization enforcement for {feature} endpoints."],
            ["TS002", f"Verify multi-currency formatting and date localization rendering in {feature}."],
            ["TS003", f"Verify browser back and forward navigation stability during multi-step flow in {feature}."],
            ["TS004", f"Verify concurrent update conflict resolution between Admin and User in {feature}."],
            ["TS005", f"Verify complete end-to-end flow of {feature} from record creation to file export."],
            ["TS006", f"Verify client-side memory usage and garbage collection during prolonged use of {feature}."],
            ["TS007", f"Verify load balancer failover transition when primary server node drops in {feature}."],
            ["TS008", f"Verify validation focus placement on first invalid input field in {feature}."],
            ["TS009", f"Verify parameter tampering protection on hidden form controls in {feature}."],
            ["TS010", f"Verify session restoration state after browser tab restart for {feature}."],
            ["TS011", f"Verify right-to-left (RTL) language text alignment and UI flip for {feature}."],
            ["TS012", f"Verify exact maximum permitted payload file upload size for {feature}."],
            ["TS013", f"Verify tooltip display on mouse hover over interactive icons in {feature}."],
            ["TS014", f"Verify malformed JSON payload rejection with HTTP 400 Bad Request in {feature}."],
            ["TS015", f"Verify asset compression Brotli and Gzip headers on static JS modules in {feature}."],
            ["TS016", f"Verify responsive breakpoint behavior at 768px for {feature} navigation drawer."],
            ["TS017", f"Verify clear success confirmation Toast notification dismissal after 5 seconds in {feature}."],
            ["TS018", f"Verify database connection pool auto-recovery after network hiccup in {feature}."],
            ["TS019", f"Verify password masking toggle icon in {feature} authentication fields."],
            ["TS020", f"Verify automatic session logout after 15 minutes of inactivity in {feature}."],
            ["TS021", f"Verify print stylesheet formatting hides navigation controls for {feature} print preview."],
            ["TS022", f"Verify clickjacking defense with X-Frame-Options SAMEORIGIN header on {feature}."],
            ["TS023", f"Verify pasting formatted rich text into plain text input stripped cleanly in {feature}."],
            ["TS024", f"Verify database isolation level prevents dirty reads during concurrent updates in {feature}."],
            ["TS025", f"Verify automated email notification dispatched upon completing action in {feature}."]
        ]

    md = "| Scenario ID | Test Scenario |\n|---|---|\n"
    for r in rows:
        md += f"| {r[0]} | {r[1]} |\n"
    return md


def generate_ai_scenarios(feature, description):
    prompt = f"""
    Generate 25 realistic and feature-specific QA test scenarios for the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Return ONLY a Markdown table with EXACTLY 2 columns:
      | Scenario ID | Test Scenario |
    - Each test scenario MUST be ONE SHORT SINGLE-LINE SENTENCE.
    - Do NOT include Category, Coverage Goal, Risk Level, Priority, Expected Result, Preconditions, or Test Steps.
    - Do NOT create long multi-line scenarios.
    - Generate at least 20 unique scenarios (prefer 20-30).
    - Cover diverse QA perspectives: happy path, negative flow, field validation, boundary conditions, error handling, network resilience, user permissions, navigation, state changes, and API/integration flows.
    - Do NOT repeat or rephrase scenarios.
    - Do NOT add introductory, explanatory, or closing text.

    Return ONLY the Markdown table.
    """
    if client:
        try:
            print("AI SCENARIOS GENERATION STARTED")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                print("AI SCENARIOS RESPONSE RECEIVED")
                return normalize_markdown_table_ids(resp.text, prefix="TS")
        except Exception as e:
            print("AI SCENARIOS FAILED:", e)

    return get_fallback_scenarios(feature, description)


def generate_ai_scenarios_regenerated(feature, description, previous_content):
    prompt = f"""
    Generate 25 completely NEW and DIFFERENT QA test scenarios for the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Do NOT repeat any scenario from the previous result provided below.
    - Do NOT merely rephrase existing scenarios.
    - Return ONLY a Markdown table with EXACTLY 2 columns:
      | Scenario ID | Test Scenario |
    - Each test scenario MUST be ONE SHORT SINGLE-LINE SENTENCE.
    - Do NOT include Category, Coverage Goal, Risk Level, Priority, Expected Result, Preconditions, or Test Steps.
    - Generate at least 20 unique scenarios (prefer 20-30).
    - Cover diverse QA perspectives: happy path, negative flow, validation, boundary conditions, error handling, network drops, role permissions, navigation, state changes, and third-party integrations.
    - Do NOT add introductory or closing text.

    PREVIOUS RESULT CONTENT TO AVOID:
    {previous_content[:1500]}

    Return ONLY the Markdown table.
    """
    if client:
        try:
            print("AI REGENERATION STARTED - SCENARIOS")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                return normalize_markdown_table_ids(resp.text, prefix="TS")
        except Exception as e:
            print("AI SCENARIO REGENERATION FAILED:", e)

    return get_fallback_scenarios_regenerated(feature, description)


def generate_ai_bug_report(feature, description):
    prompt = f"""
    Generate a realistic, detailed QA Bug Report for a defect found in the following feature.
    Feature: {feature}
    Description: {description}

    Format output clearly in Markdown with headers for:
    - Bug ID & Summary
    - Severity & Priority
    - Environment Details
    - Pre-conditions
    - Steps to Reproduce
    - Expected Result
    - Actual Result
    - Suggested Fix / Workaround
    """
    if client:
        try:
            print("AI BUG REPORT GENERATION STARTED")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                print("AI BUG REPORT RESPONSE RECEIVED")
                return resp.text
        except Exception as e:
            print("AI BUG REPORT FAILED:", e)

    fallback = f"""
### 🐛 Bug Report: Defect in {feature}

**Bug ID:** BUG-2026-001  
**Severity:** Major  
**Priority:** High  
**Environment:** Chrome / Windows 11 / Staging DB  

---

### Summary
Unexpected error occurs when submitting valid inputs in the `{feature}` module.

### Pre-conditions
1. User must be logged in as a verified QA user.
2. `{feature}` module must be accessible from dashboard.

### Steps to Reproduce
1. Navigate to **{feature}**.
2. Fill in description: *"{description}"*.
3. Click on the **Submit / Process** button.
4. Observe the system response.

### Expected Result
System processes the request successfully and displays a success confirmation message.

### Actual Result
System displays an HTTP 500 Internal Server Error message and form input is lost.

### Workaround
Refresh the page and re-enter inputs using short plain text strings.
"""
    return fallback


def generate_ai_bug_report_regenerated(feature, description, previous_content):
    prompt = f"""
    Generate a completely DIFFERENT valid QA Bug Report analysis for a defect in the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Do NOT repeat the exact defect, reproduction steps, or failure condition from the previous bug report below.
    - Vary the defect angle (e.g. state management error, concurrency race condition, validation bypass, memory leak, UI rendering glitch, or API latency failure).
    - Keep the bug realistic and relevant to the same feature and description.

    PREVIOUS BUG REPORT TO AVOID:
    {previous_content[:1500]}

    Format output clearly in Markdown with headers for:
    - Bug ID & Summary
    - Severity & Priority
    - Environment Details
    - Pre-conditions
    - Steps to Reproduce
    - Expected Result
    - Actual Result
    - Suggested Fix / Workaround
    """
    if client:
        try:
            print("AI REGENERATION STARTED - BUG REPORT")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                return resp.text
        except Exception as e:
            print("AI BUG REPORT REGENERATION FAILED:", e)

    v = random.choice([1, 2, 3])
    if v == 1:
        return f"""
### 🐛 Bug Report: Concurrent State Race Condition in {feature}

**Bug ID:** BUG-2026-002  
**Severity:** Critical  
**Priority:** High  
**Environment:** iOS Safari / Production API Endpoint  

---

### Summary
Data mutation race condition causes duplicate database entries when rapid double-clicking occurs in `{feature}`.

### Pre-conditions
1. User logged in with standard privileges.
2. Network throttled to 3G speed in DevTools.

### Steps to Reproduce
1. Open `{feature}` page.
2. Enter required payload data.
3. Rapidly double-tap the Submit button within 200ms.

### Expected Result
Submit button is disabled after initial click; single database transaction is initiated.

### Actual Result
Two simultaneous POST requests are dispatched resulting in duplicate records with duplicate IDs.

### Workaround
Implement front-end debouncing and database unique constraint verification.
"""
    elif v == 2:
        return f"""
### 🐛 Bug Report: Memory Leak and Unhandled Exception in {feature}

**Bug ID:** BUG-2026-003  
**Severity:** High  
**Priority:** High  
**Environment:** Android Chrome 124 / Staging Cluster  

---

### Summary
Heap memory allocation spikes continuously during bulk data processing in `{feature}`, crashing browser tab.

### Pre-conditions
1. Active session in `{feature}` module.
2. Dataset size exceeding 5,000 array items.

### Steps to Reproduce
1. Open `{feature}` module.
2. Trigger batch processing action for dataset.
3. Monitor Chrome Task Manager memory footprint.

### Expected Result
Processing executes in web worker thread; memory released after garbage collection.

### Actual Result
Browser tab freezes as JavaScript heap memory reaches 1.8GB limit.

### Workaround
Paginate payload processing into chunks of 500 items.
"""
    else:
        return f"""
### 🐛 Bug Report: Authorization Bypass via Parameter Tampering in {feature}

**Bug ID:** BUG-2026-004  
**Severity:** Critical  
**Priority:** Highest  
**Environment:** Chrome 124 / Staging Environment  

---

### Summary
Modifying `account_id` integer in hidden HTTP request header grants unauthorized edit access to restricted `{feature}` records.

### Pre-conditions
1. Logged in as low-privilege User A.
2. Account ID of Target User B identified.

### Steps to Reproduce
1. Intercept `POST /update_{feature}` request using Burp Suite proxy.
2. Change `X-Account-ID` header from `1001` to `1002`.
3. Forward request and observe API response.

### Expected Result
HTTP 403 Forbidden returned; authorization layer enforces ownership checks.

### Actual Result
HTTP 200 OK returned and target account data modified successfully.

### Workaround
Enforce server-side JWT claim verification on controller endpoint layer.
"""


def generate_ai_checklist(feature, description):
    prompt = f"""
    Generate a comprehensive QA Testing Checklist for the following feature.
    Feature: {feature}
    Description: {description}

    Return ONLY a Markdown table with exact columns:
    | Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
    """
    if client:
        try:
            print("AI CHECKLIST GENERATION STARTED")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                print("AI CHECKLIST RESPONSE RECEIVED")
                return resp.text
        except Exception as e:
            print("AI CHECKLIST FAILED:", e)

    fallback = f"""
| Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
|---|---|---|---|---|
| CHK_001 | Pre-conditions | Confirm user credentials and permissions for {feature}. | High | Pending |
| CHK_002 | Functional | Verify all input fields accept expected valid values. | High | Pending |
| CHK_003 | Functional | Verify clear error messaging on invalid input formats. | High | Pending |
| CHK_004 | UI / UX | Verify alignment, font typography, and button states. | Medium | Pending |
| CHK_005 | Security | Ensure sensitive inputs are masked and not exposed in URLs. | High | Pending |
| CHK_006 | Cross-Browser | Verify layout rendering on Chrome, Firefox, Safari, and Edge. | Medium | Pending |
| CHK_007 | Responsiveness | Verify display scaling on Mobile (iOS/Android) and Tablets. | Medium | Pending |
| CHK_008 | Performance | Ensure page load time is under 2.5 seconds. | Low | Pending |
"""
    return fallback


def generate_ai_checklist_regenerated(feature, description, previous_content):
    prompt = f"""
    Generate a completely NEW QA Testing Checklist for the following feature.

    Feature: {feature}
    Description: {description}

    CRITICAL INSTRUCTIONS:
    - Avoid repeating the same checklist items from the previous result below.
    - Focus on distinct inspection checkpoints (e.g. data sanitization, accessibility compliance, rate limiting, session persistence, browser caching, and responsive breakpoints).

    PREVIOUS CHECKLIST TO AVOID:
    {previous_content[:1500]}

    Return ONLY a Markdown table with exact columns:
    | Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
    """
    if client:
        try:
            print("AI REGENERATION STARTED - CHECKLIST")
            resp = client.models.generate_content(model="gemini-3.5-flash", contents=prompt)
            if resp.text:
                return resp.text
        except Exception as e:
            print("AI CHECKLIST REGENERATION FAILED:", e)

    v = random.choice([1, 2, 3])
    if v == 1:
        return f"""
| Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
|---|---|---|---|---|
| CHK_101 | Data Privacy | Verify compliance with GDPR/CCPA data export standards for {feature}. | High | Pending |
| CHK_102 | Localization | Verify date formatting, currency symbols, and UTF-8 string rendering. | Medium | Pending |
| CHK_103 | Session Security | Verify JWT token invalidation upon user logout or password reset. | High | Pending |
| CHK_104 | API Rate Limits | Verify API Gateway throttle limit response headers (X-RateLimit-Remaining). | Medium | Pending |
| CHK_105 | Input Sanitization | Verify HTML entity encoding on user text inputs to prevent XSS in {feature}. | High | Pending |
| CHK_106 | Error Handling | Verify user-friendly error messages display instead of raw 500 stack traces. | High | Pending |
| CHK_107 | Cross-Browser | Verify layout rendering on Chrome 124, Firefox ESR, Safari 17, and Edge. | Medium | Pending |
| CHK_108 | Mobile Touch | Verify tap targets exceed 48x48px minimum touch target guidelines. | Medium | Pending |
| CHK_109 | Network Resilience | Verify offline PWA caching headers for static assets in {feature}. | Low | Pending |
| CHK_110 | Performance | Verify page First Contentful Paint (FCP) is under 1.2 seconds. | High | Pending |
| CHK_111 | Accessibility | Verify all image controls have descriptive alt text tags. | Medium | Pending |
| CHK_112 | Database | Verify foreign key cascading constraints function cleanly on delete. | High | Pending |
"""
    elif v == 2:
        return f"""
| Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
|---|---|---|---|---|
| CHK_201 | Accessibility | Verify screen reader ARIA live region updates for dynamic alerts in {feature}. | High | Pending |
| CHK_202 | Offline Resilience | Verify PWA service worker offline caching for {feature} assets. | Medium | Pending |
| CHK_203 | CORS Policy | Verify Access-Control-Allow-Origin headers on {feature} API routes. | High | Pending |
| CHK_204 | Content Security | Verify CSP policy blocks inline script execution in {feature}. | High | Pending |
| CHK_205 | Form Validation | Verify mandatory field indicators (*) are clearly visible. | High | Pending |
| CHK_206 | Keyboard Navigation | Verify logical tab order across form controls without focus traps. | Medium | Pending |
| CHK_207 | Concurrency | Verify double-click prevention on submit control for {feature}. | High | Pending |
| CHK_208 | Memory Footprint | Verify memory consumption stays under 150MB during bulk tasks. | Medium | Pending |
| CHK_209 | Audit Trail | Verify user action events logged to audit table with timestamp. | Medium | Pending |
| CHK_210 | Print Preview | Verify print CSS hides navigation bars and extra controls cleanly. | Low | Pending |
| CHK_211 | Session Management | Verify automatic logout after 15 minutes of inactivity. | High | Pending |
| CHK_212 | Data Format | Verify JSON API responses format numbers and dates consistently. | Medium | Pending |
"""
    else:
        return f"""
| Item ID | Testing Category | Inspection Checkpoint | Priority | Verification Status |
|---|---|---|---|---|
| CHK_301 | Auth Scope | Verify authorization check blocks non-admin users from admin controls. | High | Pending |
| CHK_302 | Cache Invalidation | Verify browser cache cleared upon data update in {feature}. | Medium | Pending |
| CHK_303 | Mobile Viewport | Verify layout renders without horizontal scrollbar at 360px width. | High | Pending |
| CHK_304 | File Upload | Verify file extension whitelist blocks dangerous .exe and .sh files. | High | Pending |
| CHK_305 | Dark Mode | Verify text contrast ratio remains readable in Dark Mode theme. | Medium | Pending |
| CHK_306 | SQL Injection | Verify parameterized SQL queries prevent SQL injection in {feature}. | High | Pending |
| CHK_307 | HTTP Security | Verify HSTS, X-Content-Type-Options, and X-Frame-Options headers set. | High | Pending |
| CHK_308 | Auto-Complete | Verify sensitive input fields disable browser autocomplete. | Medium | Pending |
| CHK_309 | Webhook Reliability | Verify failed webhook delivery retried up to 5 times. | Medium | Pending |
| CHK_310 | Toast Notifications | Verify success notifications auto-dismiss after 4 seconds. | Low | Pending |
| CHK_311 | Data Sanitization | Verify stripping of zero-width space characters from text inputs. | Medium | Pending |
| CHK_312 | Performance | Verify DOM node count remains under 1,500 elements for performance. | Low | Pending |
"""


def regenerate_ai_content(report_type, feature, description, previous_content):
    max_attempts = 3
    best_candidate = None
    lowest_similarity = 1.0

    prefix_map = {
        "Test Case": "TC",
        "Test Scenario": "TS",
        "Bug Report": "BUG",
        "QA Checklist": "CK"
    }
    prefix = prefix_map.get(report_type, "TC")

    for attempt in range(1, max_attempts + 1):
        print(f"REGENERATION ATTEMPT {attempt} for {report_type}")

        if report_type == "Test Case":
            candidate = generate_ai_test_cases_regenerated(feature, description, previous_content)
        elif report_type == "Test Scenario":
            candidate = generate_ai_scenarios_regenerated(feature, description, previous_content)
        elif report_type == "Bug Report":
            candidate = generate_ai_bug_report_regenerated(feature, description, previous_content)
        elif report_type == "QA Checklist":
            candidate = generate_ai_checklist_regenerated(feature, description, previous_content)
        else:
            candidate = generate_ai_test_cases_regenerated(feature, description, previous_content)

        candidate = normalize_markdown_table_ids(candidate, prefix=prefix)

        sim = calculate_similarity(previous_content, candidate)
        print(f"Attempt {attempt} similarity with previous content: {sim:.2f}")

        if sim < lowest_similarity:
            lowest_similarity = sim
            best_candidate = candidate

        if sim <= 0.45:
            print(f"Regeneration successful on attempt {attempt} with low similarity ({sim:.2f})")
            return candidate

    print(f"Max attempts reached ({max_attempts}). Returning best candidate with similarity {lowest_similarity:.2f}")
    final_output = best_candidate if best_candidate else get_fallback_test_cases_regenerated(feature, description)
    return normalize_markdown_table_ids(final_output, prefix=prefix)


# -----------------------------
# Dashboard & Page Routes
# -----------------------------
@app.route("/")
def home():
    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) FROM reports")
    total_reports = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM reports WHERE type='Test Case'")
    test_cases_count = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM reports WHERE type='Test Scenario'")
    scenarios_count = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM reports WHERE type='Bug Report'")
    bugs_count = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM reports WHERE type='QA Checklist'")
    checklists_count = cursor.fetchone()[0]

    cursor.execute("SELECT * FROM reports ORDER BY id DESC LIMIT 5")
    recent_reports = cursor.fetchall()

    conn.close()

    return render_template(
        "index.html",
        total_reports=total_reports,
        test_cases_count=test_cases_count,
        scenarios_count=scenarios_count,
        bugs_count=bugs_count,
        checklists_count=checklists_count,
        recent_reports=recent_reports
    )


@app.route("/scenario", methods=["GET"])
def scenario_page():
    return render_template("scenario.html")


@app.route("/bug", methods=["GET"])
def bug_page():
    return render_template("bug.html")


@app.route("/checklist", methods=["GET"])
def checklist_page():
    return render_template("checklist.html")


# -----------------------------
# AI Action POST Routes
# -----------------------------
@app.route("/generate", methods=["POST"])
def generate():
    feature = request.form["feature"]
    description = request.form["description"]

    content = generate_ai_test_cases(feature, description)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO reports (type, feature, description, content, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
        ("Test Case", feature, description, content)
    )
    report_id = cursor.lastrowid
    conn.commit()
    conn.close()

    content_html = markdown.markdown(content, extensions=['tables'])

    return render_template(
        "result.html",
        report_id=report_id,
        report_type="Test Case",
        feature=feature,
        description=description,
        content=content,
        ai_response=content_html
    )


@app.route("/generate_scenario", methods=["POST"])
def generate_scenario():
    feature = request.form["feature"]
    description = request.form["description"]

    content = generate_ai_scenarios(feature, description)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO reports (type, feature, description, content, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
        ("Test Scenario", feature, description, content)
    )
    report_id = cursor.lastrowid
    conn.commit()
    conn.close()

    content_html = markdown.markdown(content, extensions=['tables'])

    return render_template(
        "scenario_result.html",
        report_id=report_id,
        report_type="Test Scenario",
        feature=feature,
        description=description,
        content=content,
        ai_response=content_html
    )


@app.route("/generate_bug", methods=["POST"])
def generate_bug():
    feature = request.form["feature"]
    description = request.form["description"]

    content = generate_ai_bug_report(feature, description)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO reports (type, feature, description, content, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
        ("Bug Report", feature, description, content)
    )
    report_id = cursor.lastrowid
    conn.commit()
    conn.close()

    content_html = markdown.markdown(content, extensions=['tables'])

    return render_template(
        "bug_result.html",
        report_id=report_id,
        report_type="Bug Report",
        feature=feature,
        description=description,
        content=content,
        ai_response=content_html
    )


@app.route("/generate_checklist", methods=["POST"])
def generate_checklist():
    feature = request.form["feature"]
    description = request.form["description"]

    content = generate_ai_checklist(feature, description)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO reports (type, feature, description, content, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
        ("QA Checklist", feature, description, content)
    )
    report_id = cursor.lastrowid
    conn.commit()
    conn.close()

    content_html = markdown.markdown(content, extensions=['tables'])

    return render_template(
        "checklist_result.html",
        report_id=report_id,
        report_type="QA Checklist",
        feature=feature,
        description=description,
        content=content,
        ai_response=content_html
    )


# -----------------------------
# Regeneration Route (Guarantees Feature Name Retention & Direct Processing)
# -----------------------------
@app.route("/regenerate", methods=["POST"])
def regenerate():
    report_id = request.form.get("report_id")
    feature = request.form.get("feature", "").strip()
    description = request.form.get("description", "").strip()
    report_type = request.form.get("report_type", "Test Case").strip()
    previous_content = request.form.get("previous_content", "")

    conn = get_db()
    cursor = conn.cursor()

    # Guarantee Feature Name & Description match existing report in DB
    if report_id and str(report_id).isdigit():
        try:
            cursor.execute("SELECT * FROM reports WHERE id = ?", (report_id,))
            existing = cursor.fetchone()
            if existing:
                feature = existing["feature"]
                description = existing["description"]
                report_type = existing["type"]
                if not previous_content:
                    previous_content = existing["content"] or ""
        except Exception as e:
            print("DB lookup warning during regenerate:", e)

    if not feature:
        feature = "QA Feature"

    # Execute AI regeneration loop with similarity check
    new_content = regenerate_ai_content(report_type, feature, description, previous_content)

    # Update existing entry & set real local timestamp
    try:
        if report_id and str(report_id).isdigit():
            cursor.execute(
                "UPDATE reports SET content = ?, created_at = datetime('now', 'localtime') WHERE id = ?",
                (new_content, report_id)
            )
        else:
            cursor.execute(
                "SELECT id FROM reports WHERE type = ? AND feature = ? AND description = ? ORDER BY id DESC LIMIT 1",
                (report_type, feature, description)
            )
            row = cursor.fetchone()
            if row:
                report_id = row[0]
                cursor.execute(
                    "UPDATE reports SET content = ?, created_at = datetime('now', 'localtime') WHERE id = ?",
                    (new_content, report_id)
                )
            else:
                cursor.execute(
                    "INSERT INTO reports (type, feature, description, content, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
                    (report_type, feature, description, new_content)
                )
                report_id = cursor.lastrowid

        conn.commit()
    except Exception as db_err:
        print("Database write warning on serverless environment:", db_err)
    finally:
        try:
            conn.close()
        except Exception:
            pass

    content_html = markdown.markdown(new_content, extensions=['tables'])

    template_map = {
        "Test Case": "result.html",
        "Test Scenario": "scenario_result.html",
        "Bug Report": "bug_result.html",
        "QA Checklist": "checklist_result.html"
    }
    target_template = template_map.get(report_type, "result.html")

    return render_template(
        target_template,
        report_id=report_id,
        report_type=report_type,
        feature=feature,
        description=description,
        content=new_content,
        ai_response=content_html
    )


# -----------------------------
# History & Search Routes
# -----------------------------
@app.route("/history")
@app.route("/search")
def history():
    query = request.args.get("q", "").strip()
    type_filter = request.args.get("type", "").strip()

    conn = get_db()
    cursor = conn.cursor()

    sql = "SELECT * FROM reports WHERE 1=1"
    params = []

    if query:
        sql += " AND (feature LIKE ? OR description LIKE ? OR content LIKE ?)"
        params.extend([f"%{query}%", f"%{query}%", f"%{query}%"])

    if type_filter:
        sql += " AND type = ?"
        params.append(type_filter)

    # History order DESC
    sql += " ORDER BY id DESC"

    cursor.execute(sql, params)
    reports = cursor.fetchall()
    conn.close()

    return render_template("history.html", reports=reports, query=query, type_filter=type_filter)


@app.route("/view/<int:id>")
def view_report(id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM reports WHERE id = ?", (id,))
    report = cursor.fetchone()
    conn.close()

    if not report:
        flash("Report not found.")
        return redirect(url_for("history"))

    content_html = markdown.markdown(report["content"] or "", extensions=['tables'])

    return render_template("view.html", report=report, content_html=content_html)


@app.route("/delete/<int:id>", methods=["GET", "POST"])
def delete_report(id):
    conn = get_db()
    cursor = conn.cursor()

    if request.method == "POST":
        cursor.execute("DELETE FROM reports WHERE id = ?", (id,))
        conn.commit()
        conn.close()
        return redirect(url_for("history"))

    cursor.execute("SELECT * FROM reports WHERE id = ?", (id,))
    report = cursor.fetchone()
    conn.close()

    if not report:
        flash("Report not found.")
        return redirect(url_for("history"))

    return render_template("delete_confirm.html", report=report)


# -----------------------------
# Export Routes (PDF & Excel)
# -----------------------------
@app.route("/download_pdf", methods=["POST"])
def download_pdf():
    feature = request.form.get("feature", "QA Feature")
    description = request.form.get("description", "No description provided.")
    content = request.form.get("content", "")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("<b>QA Productivity Tool - Quality Assurance Report</b>", styles["Title"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"<b>Feature:</b> {feature}", styles["Heading2"]))
    story.append(Paragraph(f"<b>Description:</b> {description}", styles["BodyText"]))
    story.append(Spacer(1, 14))

    table_rows = parse_markdown_table_data(content)
    if table_rows:
        t = Table(table_rows, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#007bff')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.grey)
        ]))
        story.append(t)
    else:
        html_text = markdown.markdown(content)
        clean_text = re.sub(r'</?(div|table|thead|tbody|tr|th|td)[^>]*>', '', html_text)
        story.append(Paragraph(clean_text, styles["BodyText"]))

    doc.build(story)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="QA_Genie_Report.pdf",
        mimetype="application/pdf"
    )


@app.route("/download_excel", methods=["POST"])
def download_excel():
    feature = request.form.get("feature", "QA Feature")
    description = request.form.get("description", "No description provided.")
    content = request.form.get("content", "")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QA Productivity Tool Report"

    sheet.append(["QA Productivity Tool - Quality Assurance Report"])
    sheet.append(["Feature", feature])
    sheet.append(["Description", description])
    sheet.append([])

    table_rows = parse_markdown_table_data(content)
    if table_rows:
        for row in table_rows:
            sheet.append(row)
    else:
        lines = content.splitlines()
        for line in lines:
            if line.strip():
                sheet.append([line.strip()])

    excel = io.BytesIO()
    workbook.save(excel)
    excel.seek(0)

    return send_file(
        excel,
        as_attachment=True,
        download_name="QA_Genie_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# -----------------------------
# Run App
# -----------------------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    app.run(debug=True, host="0.0.0.0", port=port)